"""Generate editable defense slides/report and a printable HTML deck from one source."""
from pathlib import Path
import json, shutil, html
from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches, Pt as PPTPt
from pptx.dml.color import RGBColor
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'deliverables/答辩材料'; OUT.mkdir(parents=True,exist_ok=True)
assets=OUT/'images';assets.mkdir(exist_ok=True)
for name,src in {'overview':'.runtime/ui-refactor-v4/final-overview-1920.png','stations':'.runtime/ui-refactor-v4/final-stations-1920.png','orders':'.runtime/ui-refactor-v4/final-orders-1920.png','app':'.runtime/app-reference/final-home.png','charge':'.runtime/app-reference/final-charge.png','stats':'.runtime/app-reference/final-stats.png','admin':'.runtime/admin-zh/dashboard.png','admin-station':'.runtime/admin-zh/station.png','admin-trips':'.runtime/admin-zh/trips.png','admin-history':'.runtime/admin-zh/history.png','forecast':'.runtime/forecast/admin-forecast.png','qt-user':'.runtime/qt-webmap/user/home.png','qt-admin':'.runtime/qt-webmap/admin/dashboard.png'}.items():
 if (ROOT/src).exists():shutil.copy2(ROOT/src,assets/(name+'.png'))
slides=[
('东软电动汽车充电桩\n应用管理平台',['第十一组 · 项目答辩','移动端充电服务 / 三屏运营可视化 / 服务端业务闭环','成员与指导教师：待填写 · 交付日期：2026-09-07'], '今天汇报充电桩应用管理平台。项目围绕用户找桩、预约、充电和结算建立业务闭环，同时提供三类运营大屏。我们将区分真实业务、历史数据和视觉演示，重点说明数据一致性与验证证据。',None),
('需求与交付边界',['用户：查找站点、筛选电桩、预约启停、钱包与历史订单','运营：维护用户/站点/电桩，查看订单、统计和操作日志','展示：APP 六个核心界面与总览、站点态势、订单分析大屏','新 Electra 全权限管理界面：五页面与管理抽屉已实现'], '需求既包括移动端交互，也包括管理与运营展示。独立 Electra 后台已接入现有服务，可通过页面管理用户、站点、电桩、订单与资金。本次新增验收将页面操作和数据库状态连接起来，区分模拟展示与真实业务变更。',None),
('系统架构：数据库决定业务状态',['Qt 用户端 / Qt 管理端 / Web → FastAPI REST 与 WebSocket','AccountService / ChargingService / AdminService → asyncpg','PostgreSQL：账户、站点、电桩、订单、账本、日志、事件','后台 reconcile 周期校准：超时取消、余额耗尽结算'], '采用模块化单体，避免课程项目过早拆分微服务。前端发出命令，服务端校验身份并在数据库事务中更新状态；WebSocket 用于推送变化，断线后的真实状态仍以查询接口和数据库为准。',None),
('Qt 用户端：原生 C++17 实现',['Qt 6.5.3 Widgets；首页、地图、详情、充电、统计、时段、历史、账户','原生导航、电池与能量环；WebEngine 地图支持缩放与站点联动','验证码登录、钱包充值、预约启停、结算与历史订单','Web 用户端保留；地图嵌入 WebEngine / Leaflet，其余页面保持 Qt 原生'], '根据总体架构要求新增原生Qt用户端。网络访问由C++异步请求完成，车辆和地图由Qt控件与自绘构成；服务器继续决定真实电量、费用和状态。原Web页面保留，不能把Web访问等同于运行Qt程序。','qt-user'),
('Qt 管理端：五页运营工作台',['C++17 / Qt Widgets；Qt Charts 只链接管理端','运营总览、站点导航、充电订单、操作审计、负荷预测','订单路线示意、审计时间线与热力图；管理表单接入业务接口','Ubuntu 已编译；Windows 提供构建脚本，尚未运行验证'], '管理员桌面端复用后端权限和事务边界。Qt Charts展示环图和预测折线，QPainter重绘目标环、流带与地图。管理员所有写操作通过原API确认，不直接访问数据库。用户要求跳过后续完整验收，因此不能声称Qt完整业务验收已通过。','qt-admin'),
('Qt 本地缓存与交付边界',['SQLite 仅存配置和带时间戳快照，按服务与用户隔离','JWT 内存保存；离线不能提交交易，也不离线排队','WebSocket 按会话游标重连，恢复后重新读取 REST 快照','保留 Web 大屏；本次未迁移 Vue 或新增 MLflow/ONNX'], 'SQLite不是第二套账本。客户端可在断网时查看已缓存数据，但不自行决定余额或电桩占用。Qt版本交付Linux程序、源码与Windows构建入口；完整交互和Windows运行按实际状态说明，既有Web验收不等于Qt验收。',None),
('数据模型与约束',['app_user 1—N charging_order；station 1—N charger','charger 1—N charging_order；订单关联 wallet_entry','部分唯一索引：用户与电桩各最多一个活动订单','金额采用 Decimal / numeric；行锁与事务保护结算'], '实体关系围绕订单展开。站点包含电桩，用户对电桩创建订单。活动状态的唯一索引阻止重复占用，行锁保护钱包与订单更新；金额使用十进制类型，避免二进制浮点误差。',None),
('充电闭环与资金一致性',['reserved → charging → completed','reserved → cancelled：主动取消或预约过期','预约 / 充值请求携带幂等键；结算只生成一次扣款','服务端按功率和计费时间模拟电量，余额耗尽自动结束'], '预约成功后可以开始充电，结束时统一结算。系统没有连接真实充电硬件，电量按服务端时间和功率模拟。余额耗尽时终止并将金额限制在可支付范围，重复停止请求不会再次扣款。',None),
('移动端：深紫色充电体验',['车辆首页、站点地图、充电详情、充电中、统计、时段电价','本地深圳道路数据与地图标记；电桩筛选、限额预览','真实预约需要用户登录；车辆外观与默认电量是演示'], '移动端参照用户提供的图片实现深紫色卡片、渐变按钮和车辆视觉。浏览无需登录，真正预约需要用户身份。默认车型、电池百分比和续航用于界面展示，不等同于车辆实时遥测。','app'),
('充电中与用户操作',['能量光环突出充电进度；结束充电操作固定可见','实际订单的电量、费用与状态由后端返回','保留 0.00 元与 0.0 kWh，不用演示数字覆盖零值','订单历史与钱包形成可回溯操作链'], '这里展示充电状态页。视觉效果与数据分开处理，真实订单刚开始时费用为零，我们专门验证零值不会被演示值替换。操作结束后用户可回到订单历史查看结果。','charge'),
('运营总览：高密度历史分析',['紫蓝色地球与中国轮廓，关键指标位于左侧','右侧 15 个细分图表，支持多维历史数据展示','当前业务 14 个站点 / 178 个电桩','UrbanEV 历史与平台当前业务采用独立口径'], '总览通过地球聚焦和密集图表呈现运营信息。必须区分当前平台的十四个站点、一百七十八个电桩，与导入的 UrbanEV 历史记录，不能把历史数据当作当前实时业务规模。','overview'),
('站点态势与订单分析',['站点态势：深圳行政区域聚焦与空间分布','订单分析：股票走势式主图、均线、成交量式柱图','30 / 90 / 181 天切换，三屏合计 45 个右侧小图','演示大屏默认只读管理员身份；管理写操作仍校验权限'], '三屏保持统一视觉语言，但主图承担不同任务。站点页强调空间分布，订单分析用趋势与均线呈现时间变化。默认演示令牌只有读取权限，与真实管理账号不同。','orders'),
('UrbanEV：原始文件到可视化',['固定 Git 提交下载，保存文件散列与导入元数据','GitHub data 目录 23 个原始文件完整归档','1,362 条站点记录；1,194,600 条区域小时观测','2022-09-01 至 2023-02-28：275 区域 × 4,344 小时'], '数据来自 UrbanEV 开源仓库。我们导入的是该提交的 data 目录，不包括外部网盘上另行发布的资料。二十三个文件保留原始字节并核对散列，六种指标按区域与小时对齐后形成一百一十九万余条观测。',None),
('管理员 Dashboard：视觉与操作',['三上二下卡片：透视车、半圆目标、路线、流带与站点','运营总览 / 站点导航 / 充电订单 / 操作审计 / 负荷预测','管理抽屉覆盖用户、站点、电桩、订单、资金与设置','默认全权限模拟账号；真实变更入库并记录审计'], '管理员后台按指定参考图复刻布局、色彩和关键图形，使用真实站点列表。深层管理通过抽屉和表单打开。独立开发开关开启时默认签发全权限管理员会话，其他环境默认关闭。','admin'),
('站点与全权限管理',['站点导航：竖向透视车辆与大幅路线示意','用户资料/冻结、站点及设备增删改、设备故障恢复','代预约、启动、取消、结算；调账、账本及部分退款','模拟路线不提供真实导航；车辆与目标不作计费依据'], '站点页延续参考图的空间构图。管理员可以代用户完成完整充电流程，但不能绕过余额、设备占用和历史保留约束。路线与车辆是视觉演示，订单和账本来自数据库。','admin-station'),
('充电订单：充电旅程与订单分析',['路线示意、订单状态环图、每日能量曲线','带站点缩略图与连接示意的订单列表','支持状态筛选、搜索、详情、启动、取消、结束与退款','摘要取已加载订单；设备可用率取当前设备状态'], '这一页先由图像生成工具产出与首页一致的视觉概念，再提取路线卡、环图和订单条目并实现为真实网页。所有数字由服务端记录计算，路线仅为示意，空白日期显示零。','admin-trips'),
('操作审计：从日志表到审计时间线',['按日期分组的彩色事件节点，点击联动详情','操作构成环图、UTC 星期/小时热力图、14天活动柱图','类型筛选、关键词搜索、完整记录查看及 CSV 导出','审计数据不补造；无事件时保留真实空状态'], 'History 用可交互时间线呈现审计记录。右侧图表展示实际操作构成和时间分布，选中事件后同步显示详情。界面延续紫黑卡片与粉紫、金色、青色点缀，同时保留完整操作标识便于核对。','admin-history'),
('负荷预测：历史实验已接入',['管理员第五页：全网与 275 个区域的未来 24 小时预测','最近48小时、预测曲线、经验误差范围与历史高负荷阈值','直接读取数据库119万条区域小时记录，训练岭回归','数据截止2023-02-28：预测2023-03-01，非实时结果'], '负荷预测已真正训练并接入管理员界面。输入是数据库的区域小时估算电量，不是演示订单或站点实测功率。页面同时显示数据日期、模型比较和范围说明。全网和区域分别建模，训练脚本与系数产物可复现。','forecast'),
('模型评估：先与简单基线比较',['时间划分：1月17日前训练，1月18—31日验证，2月测试','全网测试 WAPE：昨日6.260%、上周10.257%、岭回归6.055%','仅按验证期选模型：全网仍采用昨日基线，99区域采用岭回归','提升有限；六个月估算数据不代表当前实测效果'], '每个预测起点只使用之前的历史。测试期逐日预测24小时，参数固定；模型选择只看验证集，不能因测试集岭回归略好就改选。最后才使用全部历史重新拟合用于下一日展示。WAPE不是准确率。源数据可能有非因果预处理，误差范围是经验范围而非严格置信保证。',None),
('管理资金：权限完整、约束保留',['资金操作：用户行锁、幂等键与钱包账本','重复退款只记一次；累计退款不超过结算金额','活动订单禁止改设备参数、负向调账；历史数据保留','日志写入失败时，订单与电桩更新一起回滚'], '全权限表示管理员可发起所有业务操作，不表示可以破坏账务约束。我们通过并发重复调账、重复退款和超额退款测试验证边界。通过制造日志外键失败，确认日志与业务状态处在同一事务。',None),
('核心服务：既有验证证据',['后端 smoke：认证、充值幂等、设备管理、充电与 WebSocket','隔离模式验证：两用户争抢同一电桩，仅一个成功','验证余额耗尽结算、重复结束不重复扣款、预约过期释放','数据导入校验、数据库实际恢复、浏览器交互另附日志'], '测试不只检查页面能打开。我们用隔离数据库模式同时发送预约请求，验证只允许一个成功；并通过推进服务端时间验证余额耗尽和过期释放。所有结论对应交付包内日志，浏览器夹具测试与真实后端测试分别记录。',None),
('部署与交付',['Python 3.12 + PostgreSQL 14；wheel / 源代码两种安装方式','独立 tmux 会话 assignment-public：web 与 funnel 两个窗口','Tailscale Funnel 公网 HTTPS 入口；与本次对话进程分离','完整私有备份 + 整理版源码数据包 + 可恢复数据库备份'], '当前服务与 Funnel 运行在独立 tmux 会话中，终端断开不影响运行。但机器重启后仍需启动脚本，公网可用性也取决于主机与 Tailscale 服务。本次打包同时提供私有环境快照和便于迁移的整理版，恢复以逻辑数据库备份为准。',None),
('限制与后续工作',['开发 OTP、模拟充电与演示数据不能等同于商用接桩','无真实支付、硬件遥测或已验证的高并发性能结论','待完成：正式认证、支付与设备协议接入','后续：峰谷策略、可观测性、负载测试与部署自动化'], '目前项目完成课程演示所需的业务闭环与可视化，但不是商用充电平台。我们没有接入真实支付和设备协议，也没有用百万条历史数据证明系统具备百万用户并发能力。下一步推进真实认证、支付和设备协议，并开展负载与稳定性测试。',None),
('演示路径与讨论',['1 分钟：APP 首页 → 站点 → 电桩详情','2 分钟：用户登录 → 充值 → 预约 → 启停 → 历史','1 分钟：三屏切换 → 站点聚焦 → 趋势范围','1 分钟：管理员代预约、调账/退款与审计日志'], '最后按这条五分钟路径演示。操作前确认公网和本地健康检查，账户不足时先充值；若网络异常，直接切换本地地址与截图材料。欢迎老师针对事务、权限边界、数据口径与后续实现提问。',None),
]
qas=[('为什么选择模块化单体？','当前规模下便于维护事务边界和部署；模块职责清晰，之后有实际瓶颈再拆服务。'),('怎么防止两个人抢同一个电桩？','事务、行锁与活动订单部分唯一索引共同保护；并发验收是一成功一冲突。'),('一个用户能预约两台吗？','不能同时持有两个活动订单，数据库约束和服务校验共同限制。'),('为什么不用浮点数计算钱？','二进制浮点难以精确表示部分小数，系统使用 Decimal 与 PostgreSQL numeric。'),('重复点击停止会重复扣费吗？','结算在事务中检查订单状态；重复停止返回已完成结果，验收确认只有一条扣款记录。'),('余额不足怎么办？','周期校准结束订单，将扣款限制在余额内；已验证 5 元余额耗尽后余额为零。'),('预约超时如何释放？','后台校准将过期预约取消并释放电桩，相关隔离测试通过。'),('WebSocket 断线会不会丢业务？','数据库为事实来源；重新查询 REST 获取状态，推送不承担唯一账务记录。'),('这是实际充电数据吗？','平台充电过程由服务端时间和功率模拟；UrbanEV 是公开历史观测；车辆默认视觉值另属演示。'),('数据有多少？','23 个原始文件、1362 条站点记录、1194600 条区域小时观测；当前业务为14站178桩。'),('为什么观测数是1194600？','275 个区域乘以 4344 小时，六个指标对齐在每条区域小时观测中。'),('有没有导入 UrbanEV 全部资料？','固定提交 GitHub data 目录内23个文件已归档；不包含外部网盘单独分发的数据。'),('数据许可是什么？','UrbanEV 仓库标注 CC0；应保留项目来源、固定提交及文件校验信息。'),('默认管理员能修改吗？','现有大屏默认是只读 demo_admin；现有真实管理员接口可写，新管理后台可在开发环境通过独立开关自动获得全权限会话。'),('新管理员界面做到哪里？','/ui/admin.html 已实现五页面和管理抽屉；支持用户、站点、电桩、订单、调账、退款和设置。模拟地图与车辆不代表硬件遥测。'),('如何证明测试不是演示？','后端 smoke 和并发测试调用真实隔离服务；APP流程有独立浏览器夹具，日志明确区分二者。'),('百万数据意味着支持百万并发吗？','不意味着；导入规模是存储量，尚未进行足以证明高并发容量的负载测试。'),('断开终端网站会停吗？','不会因终端断开而停，web与Funnel在独立tmux会话；机器重启后需重新启动。'),('如何恢复数据库？','用 pg_restore 将 custom 格式备份恢复到新数据库；交付前已经实际恢复并对比十一张关键表行数。'),('下一步最重要的事是什么？','接入真实认证、支付和充电设备协议，完善分页交互、负载测试与可观测性。')]
qas += [('Qt 是网页套壳吗？','不是。主体由C++17和Qt Widgets构成，地图区域使用QWebEngine加载Leaflet和OSM在线底图；Qt Charts仅管理端链接，网络和本地缓存分别使用Qt Network/WebSockets/SQL。'),('Qt 重构验收完成了吗？','两个Linux程序已编译。按用户要求跳过后续完整验收，Windows未运行；不得把既有Web和后端验收结果视为Qt完整验收。'),('模型真的训练了吗？','训练了按范围独立拟合的岭回归，保留系数、输入数据散列和复现脚本，同时对比昨日与上周周期基线。详见 docs/INTEGRATION_LINUX.md。'),('为什么岭回归测试更好却选择昨日基线？','模型按验证集选择。全网验证期昨日基线更好，不能在看到测试集后改选，否则测试不再独立。'),('这能预测今天的负荷吗？','不能。当前数据截止2023年2月，页面是次日历史预测实验。需要接入更新的实测电量后再验证上线。')]
intro='# 项目答辩讲稿与技术说明\n\n东软电动汽车充电桩应用管理平台 · 第十一组\n\n成员、学号、指导教师：待填写。建议汇报 8–10 分钟，演示 5 分钟。\n\n'
report=intro
for i,(title,bullets,note,img) in enumerate(slides,1):
 report+=f'## {i}. {title.replace(chr(10), " ")}\n\n'+''.join('- '+x+'\n' for x in bullets)+'\n讲稿：'+note+'\n\n'
report+=f'## 常见问题（{len(qas)} 问）\n\n'+''.join(f'### {i}. {q}\n\n{a}\n\n' for i,(q,a) in enumerate(qas,1))
report+='## 来源与口径\n\n- UrbanEV：https://github.com/IntelligentSystemsLab/UrbanEV ，固定提交 44f2aa0c8d89f192bce00bafb0def74a21b39c68。\n- APP 视觉依据用户上传参考图；新管理后台参考：https://dribbble.com/shots/27064652-EV-Charging-Dashboard-UI-UX-EV-Charging-Station-App-UI-Design 。设计参考不代表原作者参与开发。\n- 车辆及充电光环使用 AI 生成图片；项目需求、完成度和头像图像处理说明见 docs/REQUIREMENTS.md。\n- 验证日志位于 deliverables/evidence；可复现源码脚本位于 scripts 与完整备份的 .runtime 中。\n'
(OUT/'技术说明与答辩讲稿.md').write_text(report)
doc=Document();doc.styles['Normal'].font.name='Microsoft YaHei';doc.styles['Normal'].font.size=Pt(10)
for line in report.splitlines():
 if line.startswith('# '):doc.add_heading(line[2:],0)
 elif line.startswith('## '):doc.add_heading(line[3:],1)
 elif line.startswith('### '):doc.add_heading(line[4:],2)
 elif line.startswith('- '):doc.add_paragraph(line[2:],style='List Bullet')
 elif line:doc.add_paragraph(line)
doc.save(OUT/'技术说明与答辩讲稿.docx')
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
def textbox(slide,x,y,w,h,text,size,color='FFFFFF'):
 box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=box.text_frame;tf.word_wrap=True
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Microsoft YaHei';p.font.size=PPTPt(size);p.font.color.rgb=RGBColor.from_string(color);p.space_after=PPTPt(18)
 return box
pages=[]
for i,(title,bullets,note,img) in enumerate(slides,1):
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('100A21')
 textbox(slide,.6,.35,12,.4,'ELECTRA / EV CHARGING PLATFORM · 第十一组',12,'BE9FFF')
 textbox(slide,.6,1,12,1.25,title,30)
 textbox(slide,.7,2.55,6 if img else 11.8,4,'\n'.join('• '+x for x in bullets),20,'E1D8F1')
 if img and (assets/(img+'.png')).exists():
  path=assets/(img+'.png');w,h=Image.open(path).size;scale=min(5.5/w,4.65/h);iw=w*scale;ih=h*scale
  slide.shapes.add_picture(str(path),Inches(7+(5.5-iw)/2),Inches(2.2+(4.65-ih)/2),width=Inches(iw),height=Inches(ih))
 slide.notes_slide.notes_text_frame.text=note
 textbox(slide,.7,7,12,.3,f'2026.09.07  /  交付与答辩                         {i:02d} / {len(slides):02d}',10,'9B88B4')
 picture=f'<img src="images/{img}.png">' if img else ''
 pages.append('<section><small>ELECTRA / EV CHARGING PLATFORM · 第十一组</small><h1>'+html.escape(title).replace('\n','<br>')+'</h1><article><ul>'+''.join('<li>'+html.escape(x)+'</li>' for x in bullets)+'</ul>'+picture+f'</article><footer>2026.09.07 · 项目答辩<span>{i:02d} / {len(slides):02d}</span></footer></section>')
prs.save(OUT/'答辩汇报.pptx')
font=ROOT/'src/charging_core/static/assets';fonts=list(font.rglob('*.woff2'))
# Browser can use the local CJK fonts without external requests.
fontcss=''
for f in fonts:
 if '400' in f.name or 'regular' in f.name.lower():
  shutil.copy2(f,assets/f.name);fontcss=f"@font-face{{font-family:LocalCJK;src:url('images/{f.name}')}}";break
(OUT/'答辩汇报.html').write_text('<!doctype html><meta charset="utf-8"><title>第十一组 · 项目答辩</title><style>'+fontcss+'*{box-sizing:border-box}body{margin:0;background:#100a21;color:#fff;font-family:LocalCJK,"Noto Sans CJK SC",sans-serif}section{width:1280px;height:720px;padding:36px 55px;position:relative;page-break-after:always;background:radial-gradient(ellipse at 90% 0,#34204e,#100a21 60%)}small{color:#bc9feb;letter-spacing:2px}h1{font-size:38px;line-height:1.3;margin:30px 0 20px}article{display:flex;align-items:center;gap:32px;height:450px}ul{flex:1;padding-left:26px;font-size:25px;line-height:1.65;color:#e1d8f1}li{margin-bottom:21px}img{max-width:550px;max-height:440px;object-fit:contain;border:1px solid #4c3862;border-radius:12px}footer{position:absolute;bottom:24px;left:55px;right:55px;color:#9b88b4;font-size:13px}footer span{float:right}@page{size:1280px 720px;margin:0}@media print{section{break-inside:avoid;print-color-adjust:exact;-webkit-print-color-adjust:exact}}</style>'+''.join(pages))
print('Generated administrator edition PPTX + HTML and editable DOCX/Markdown report')
